# Copyright (c) 2026, Isoft and Contributors
# MIT License. See license.txt
"""Builds a formatted .xlsx sheet out of child table rows.

Kept separate from api.py so it can be reused (reports, scheduled exports)
without going through the whitelisted endpoint.
"""

from __future__ import unicode_literals

import datetime
import re
from io import BytesIO

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import frappe
from frappe import _
from frappe.utils import cint, flt, get_datetime, getdate

# fieldtypes that carry a number Excel should treat as a number
NUMERIC_TYPES = ("Currency", "Float", "Int", "Percent")
# only these get a totals row - summing a rate or a percent is meaningless
SUMMABLE_TYPES = ("Currency", "Float", "Int")
HTML_TYPES = ("Text Editor", "HTML", "Markdown Editor", "Code")

ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")
INVALID_SHEET_CHARS_RE = re.compile(r"[\[\]:*?/\\]")

# palette lifted from the desk so the sheet looks like the app it came from
COLOR_TEXT = "1F272E"
COLOR_MUTED = "8D99A6"
COLOR_HEADER_BG = "4E5B6E"
COLOR_BORDER = "D1D8DD"
COLOR_BAND = "F7F9FA"
COLOR_TOTAL_BG = "EBEFF2"

THIN = Side(style="thin", color=COLOR_BORDER)
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BORDER = Border(
	left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color=COLOR_HEADER_BG)
)


def clean(value):
	"""openpyxl refuses control characters."""
	if isinstance(value, str):
		value = ILLEGAL_CHARACTERS_RE.sub("", value)
	return value


def strip_html(value):
	from frappe.utils.xlsxutils import handle_html

	try:
		return handle_html(value)
	except Exception:
		return frappe.utils.strip_html(value)


def get_date_format():
	fmt = frappe.db.get_single_value("System Settings", "date_format") or "yyyy-mm-dd"
	return fmt.upper()


def get_time_format():
	fmt = frappe.db.get_single_value("System Settings", "time_format") or "HH:mm:ss"
	# frappe stores HH:mm:ss / HH:mm - excel wants the same letters uppercased
	return fmt.replace("mm", "MM").upper()


def number_format(df):
	"""Excel number format string for a numeric docfield."""
	fieldtype = df.get("fieldtype")

	if fieldtype == "Int":
		return "#,##0"

	if fieldtype == "Percent":
		return "0.00%"

	precision = cint(df.get("precision"))
	if not precision:
		if fieldtype == "Currency":
			precision = cint(frappe.db.get_default("currency_precision")) or 2
		else:
			precision = cint(frappe.db.get_default("float_precision")) or 2

	return "#,##0" + ("." + "0" * precision if precision else "")


def cell_value(value, df):
	"""Convert a stored value into something Excel understands natively."""
	fieldtype = df.get("fieldtype")

	if value in (None, ""):
		# 0 is meaningful for numbers, blank is meaningful for everything else
		return 0 if fieldtype in NUMERIC_TYPES else ""

	if fieldtype == "Check":
		return _("Yes") if cint(value) else _("No")

	if fieldtype == "Percent":
		# stored as 20 for 20%, but the 0.00% format multiplies by 100
		return flt(value) / 100.0

	if fieldtype in ("Currency", "Float"):
		return flt(value)

	if fieldtype == "Int":
		return cint(value)

	if fieldtype == "Date":
		try:
			return getdate(value)
		except Exception:
			return clean(str(value))

	if fieldtype == "Datetime":
		try:
			dt = get_datetime(value)
			# excel has no timezone concept, drop microseconds for a clean cell
			return dt.replace(microsecond=0)
		except Exception:
			return clean(str(value))

	if fieldtype == "Time":
		try:
			if isinstance(value, datetime.timedelta):
				total = int(value.total_seconds())
				return datetime.time(total // 3600 % 24, total // 60 % 60, total % 60)
			return get_datetime("2000-01-01 " + str(value)).time()
		except Exception:
			return clean(str(value))

	if fieldtype in HTML_TYPES:
		return clean(strip_html(frappe.as_unicode(value)))

	return clean(frappe.as_unicode(value))


def display_width(value, df):
	"""Rough rendered width of a value, used to size columns."""
	fieldtype = df.get("fieldtype")

	if isinstance(value, (datetime.datetime,)):
		return len(get_date_format()) + len(get_time_format()) + 1
	if isinstance(value, datetime.date):
		return len(get_date_format())
	if isinstance(value, datetime.time):
		return len(get_time_format())
	if isinstance(value, (int, float)) and fieldtype in NUMERIC_TYPES:
		# thousands separators and decimals add characters the raw repr lacks
		return len("{:,.2f}".format(value)) + 1

	text = frappe.as_unicode(value or "")
	if not text:
		return 0
	# multi-line text wraps, so only the longest line matters
	return max(len(line) for line in text.split("\n"))


def style_data_cell(cell, df, date_fmt, time_fmt):
	"""Border, font, number format and alignment for one data cell."""
	fieldtype = df.get("fieldtype")

	cell.border = CELL_BORDER
	cell.font = Font(name="Calibri", size=10, color=COLOR_TEXT)

	if fieldtype in NUMERIC_TYPES:
		cell.number_format = number_format(df)
		cell.alignment = Alignment(horizontal="right", vertical="top")
	elif fieldtype == "Date":
		cell.number_format = date_fmt
		cell.alignment = Alignment(horizontal="center", vertical="top")
	elif fieldtype == "Datetime":
		cell.number_format = date_fmt + " " + time_fmt
		cell.alignment = Alignment(horizontal="center", vertical="top")
	elif fieldtype == "Time":
		cell.number_format = time_fmt
		cell.alignment = Alignment(horizontal="center", vertical="top")
	elif fieldtype == "Check":
		cell.alignment = Alignment(horizontal="center", vertical="top")
	else:
		cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def build_workbook(columns, rows, title, subtitle=None, add_totals=False, sheet_name=None):
	"""Return a BytesIO xlsx.

	:param columns: list of docfield-ish dicts, each needs `fieldname`, `label`, `fieldtype`
	:param rows: list of dicts keyed by fieldname
	:param title: heading printed in A1
	:param subtitle: small grey line under the title
	:param add_totals: append a SUM row for Currency/Float/Int columns
	"""
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = safe_sheet_name(sheet_name or title)

	last_col = len(columns)
	last_letter = get_column_letter(last_col)

	# ---- title block --------------------------------------------------
	ws["A1"] = clean(title)
	ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=COLOR_TEXT)
	ws["A2"] = clean(subtitle or "")
	ws["A2"].font = Font(name="Calibri", size=9, italic=True, color=COLOR_MUTED)
	if last_col > 1:
		ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
		ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
	ws.row_dimensions[1].height = 22
	ws.row_dimensions[3].height = 6

	# ---- header row ---------------------------------------------------
	header_row = 4
	header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
	widths = []

	for i, df in enumerate(columns, start=1):
		label = _(df.get("label") or df.get("fieldname"))
		cell = ws.cell(row=header_row, column=i, value=clean(label))
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = CELL_BORDER
		widths.append(len(label))

	ws.row_dimensions[header_row].height = 24

	# ---- data ---------------------------------------------------------
	band_fill = PatternFill("solid", fgColor=COLOR_BAND)
	date_fmt = get_date_format()
	time_fmt = get_time_format()

	for r, row in enumerate(rows):
		excel_row = header_row + 1 + r
		banded = r % 2 == 1

		for i, df in enumerate(columns, start=1):
			value = cell_value(row.get(df.get("fieldname")), df)
			cell = ws.cell(row=excel_row, column=i, value=value)
			style_data_cell(cell, df, date_fmt, time_fmt)

			if banded:
				cell.fill = band_fill

			widths[i - 1] = max(widths[i - 1], display_width(value, df))

	last_data_row = header_row + len(rows)

	# ---- totals -------------------------------------------------------
	if add_totals and rows:
		total_row = last_data_row + 1
		ws.cell(row=total_row, column=1, value=_("Total"))

		for i, df in enumerate(columns, start=1):
			cell = ws.cell(row=total_row, column=i)
			cell.font = Font(name="Calibri", bold=True, size=10, color=COLOR_TEXT)
			cell.fill = PatternFill("solid", fgColor=COLOR_TOTAL_BG)
			cell.border = TOTAL_BORDER

			if df.get("fieldtype") in SUMMABLE_TYPES and df.get("fieldname") != "idx":
				letter = get_column_letter(i)
				cell.value = "=SUM({0}{1}:{0}{2})".format(letter, header_row + 1, last_data_row)
				cell.number_format = number_format(df)
				cell.alignment = Alignment(horizontal="right")
			elif i == 1:
				cell.alignment = Alignment(horizontal="left")

	# ---- finishing ----------------------------------------------------
	for i, width in enumerate(widths, start=1):
		ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 55)

	ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
	if rows:
		ws.auto_filter.ref = "A{0}:{1}{2}".format(header_row, last_letter, last_data_row)

	ws.sheet_view.showGridLines = False
	ws.page_setup.orientation = "landscape" if last_col > 6 else "portrait"
	ws.print_title_rows = "{0}:{0}".format(header_row)

	out = BytesIO()
	wb.save(out)
	return out


def safe_sheet_name(name):
	name = INVALID_SHEET_CHARS_RE.sub(" ", frappe.as_unicode(name or "Sheet1")).strip()
	return (name or "Sheet1")[:31]


def safe_filename(name):
	name = re.sub(r"[^\w\-. ]+", "_", frappe.as_unicode(name or "export")).strip()
	return (name or "export")[:120]


# ---------------------------------------------------------------------------
# Bulk edit template: the xlsx that replaces the core CSV round trip
# ---------------------------------------------------------------------------

# row 3 carries the fieldnames and is hidden; the parser reads it back
TEMPLATE_FIELDNAME_ROW = 3
TEMPLATE_HEADER_ROW = 4
# spare rows kept formatted (and validated) so users can just keep typing
TEMPLATE_BLANK_ROWS = 50

COLOR_REQD_BG = "8D4B4B"
COLOR_NOTE = "6C7680"


def build_template_workbook(columns, rows, title, note=None, sheet_name=None):
	"""Return a BytesIO xlsx that can be edited and uploaded straight back.

	Layout: title, note, a hidden fieldname row, the labels header, then data.
	The hidden row is what makes the round trip safe - labels are translated and
	get renamed, fieldnames do not.

	:param columns: docfield-ish dicts with `fieldname`, `label`, `fieldtype`,
		optionally `reqd`, `options`, `description`
	:param rows: list of dicts keyed by fieldname, the rows currently in the grid
	"""
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = safe_sheet_name(sheet_name or title)

	last_col = len(columns)

	# ---- title block --------------------------------------------------
	ws["A1"] = clean(_("Bulk Edit {0}").format(title))
	ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=COLOR_TEXT)
	ws["A2"] = clean(
		note
		or _(
			"Edit the rows below and upload this file back with the Upload button. "
			"Do not rename or reorder the header row. Columns marked * are mandatory."
		)
	)
	ws["A2"].font = Font(name="Calibri", size=9, italic=True, color=COLOR_NOTE)
	if last_col > 1:
		ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
		ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
	ws.row_dimensions[1].height = 22

	# ---- hidden fieldname row -----------------------------------------
	for i, df in enumerate(columns, start=1):
		cell = ws.cell(row=TEMPLATE_FIELDNAME_ROW, column=i, value=df.get("fieldname"))
		cell.font = Font(name="Calibri", size=8, color=COLOR_MUTED)
	ws.row_dimensions[TEMPLATE_FIELDNAME_ROW].hidden = True

	# ---- header row ---------------------------------------------------
	header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
	reqd_fill = PatternFill("solid", fgColor=COLOR_REQD_BG)
	widths = []

	for i, df in enumerate(columns, start=1):
		label = _(df.get("label") or df.get("fieldname"))
		if df.get("reqd"):
			label += " *"

		cell = ws.cell(row=TEMPLATE_HEADER_ROW, column=i, value=clean(label))
		cell.font = header_font
		cell.fill = reqd_fill if df.get("reqd") else header_fill
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = CELL_BORDER

		hint = column_hint(df)
		if hint:
			cell.comment = Comment(hint, "ISOFT", height=110, width=260)

		widths.append(len(label))

	ws.row_dimensions[TEMPLATE_HEADER_ROW].height = 24

	# ---- data ---------------------------------------------------------
	date_fmt = get_date_format()
	time_fmt = get_time_format()
	first_data_row = TEMPLATE_HEADER_ROW + 1

	for r, row in enumerate(rows):
		excel_row = first_data_row + r
		for i, df in enumerate(columns, start=1):
			value = template_value(row.get(df.get("fieldname")), df)
			cell = ws.cell(row=excel_row, column=i, value=value)
			style_data_cell(cell, df, date_fmt, time_fmt)
			if df.get("fieldtype") == "Percent":
				# store 20, not 0.2 - the sheet has to round trip through a
				# parser that cannot tell "20%" typed as 20 from 20% as 0.2
				cell.number_format = '0.00"%"'
			widths[i - 1] = max(widths[i - 1], display_width(value, df))

	# keep the formatting going past the last row so added rows look the same
	for r in range(len(rows), len(rows) + TEMPLATE_BLANK_ROWS):
		excel_row = first_data_row + r
		for i, df in enumerate(columns, start=1):
			cell = ws.cell(row=excel_row, column=i)
			style_data_cell(cell, df, date_fmt, time_fmt)
			if df.get("fieldtype") == "Percent":
				cell.number_format = '0.00"%"' 

	last_row = first_data_row + len(rows) + TEMPLATE_BLANK_ROWS - 1
	add_validations(ws, columns, first_data_row, last_row)

	# ---- finishing ----------------------------------------------------
	for i, width in enumerate(widths, start=1):
		ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 12), 45)

	ws.freeze_panes = ws.cell(row=first_data_row, column=1)
	ws.sheet_view.showGridLines = False
	ws.page_setup.orientation = "landscape" if last_col > 6 else "portrait"
	ws.print_title_rows = "{0}:{0}".format(TEMPLATE_HEADER_ROW)

	out = BytesIO()
	wb.save(out)
	return out


def template_value(value, df):
	"""cell_value, tuned for a sheet that gets read back.

	Blank stays blank even for numbers - a 0 written into every empty numeric
	cell comes back as a real 0 on upload and quietly overrides field defaults.
	Percent keeps the number the user reads (20, not 0.2), see the format below.
	"""
	if value in (None, ""):
		return ""

	if df.get("fieldtype") == "Percent":
		return flt(value)

	return cell_value(value, df)


def column_hint(df):
	"""Tooltip pinned to the header cell, so the sheet explains itself."""
	parts = [df.get("fieldname")]

	fieldtype = df.get("fieldtype")
	if fieldtype == "Link":
		parts.append(_("Link to {0} - type the exact ID").format(_(df.get("options") or "")))
	elif fieldtype == "Check":
		parts.append(_("Yes or No"))
	elif fieldtype == "Date":
		parts.append(get_date_format())
	elif fieldtype == "Datetime":
		parts.append(get_date_format() + " " + get_time_format())
	elif fieldtype:
		parts.append(_(fieldtype))

	if df.get("reqd"):
		parts.append(_("Mandatory"))

	if df.get("read_only"):
		parts.append(_("Calculated - may be recalculated when the document is saved"))

	description = strip_html(df.get("description") or "").strip()
	if description:
		parts.append(description)

	return "\n".join(clean(p) for p in parts if p)


def add_validations(ws, columns, first_row, last_row):
	"""Dropdowns for Check and short Select columns - fewer typos on upload."""
	from openpyxl.worksheet.datavalidation import DataValidation

	for i, df in enumerate(columns, start=1):
		fieldtype = df.get("fieldtype")

		if fieldtype == "Check":
			choices = [_("Yes"), _("No")]
		elif fieldtype == "Select":
			choices = [o.strip() for o in (df.get("options") or "").split("\n") if o.strip()]
		else:
			continue

		formula = '"{0}"'.format(",".join(c.replace('"', "") for c in choices))
		# excel caps the inline list at 255 characters
		if not choices or len(formula) > 255:
			continue

		dv = DataValidation(type="list", formula1=formula, allow_blank=True)
		dv.error = _("Pick one of: {0}").format(", ".join(choices))
		dv.errorTitle = _("Invalid value")
		ws.add_data_validation(dv)
		letter = get_column_letter(i)
		dv.add("{0}{1}:{0}{2}".format(letter, first_row, last_row))


def read_spreadsheet(content, filename=None):
	"""Read an uploaded xlsx / xls / csv into a list of rows of raw cell values."""
	from frappe.utils.csvutils import read_csv_content
	from frappe.utils.xlsxutils import (
		read_xls_file_from_attached_file,
		read_xlsx_file_from_attached_file,
	)

	extension = (filename or "").rsplit(".", 1)[-1].lower()

	if extension == "csv":
		return read_csv_content(frappe.safe_decode(content))

	if extension == "xls":
		return read_xls_file_from_attached_file(content)

	# .xlsx and anything else worth a try - openpyxl fails loudly on garbage
	try:
		return read_xlsx_file_from_attached_file(fcontent=content)
	except Exception:
		if extension in ("xlsx", "xlsm"):
			raise
		return read_csv_content(frappe.safe_decode(content))
