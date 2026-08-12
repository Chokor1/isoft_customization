# Copyright (c) 2026, Isoft and Contributors
# MIT License. See license.txt
"""Builds a formatted .xlsx sheet out of child table rows.

Kept separate from api.py so it can be reused (reports, scheduled exports)
without going through the whitelisted endpoint.
"""

from __future__ import unicode_literals

import datetime
import re
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import frappe
from frappe import _
from frappe.utils import cint, flt, get_datetime, getdate

# fieldtypes that carry a number Excel should treat as a number
NUMERIC_TYPES = ("Currency", "Float", "Int", "Percent")
# only these get a totals row - summing a rate or a percent is meaningless
SUMMABLE_TYPES = ("Currency", "Float", "Int")
HTML_TYPES = ("Text Editor", "HTML", "Markdown Editor", "Code")

ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")
INVALID_SHEET_CHARS_RE = re.compile(r"[\[\]:*?/\\]")

# palette lifted from the desk so the sheet looks like the app it came from
COLOR_TEXT = "1F272E"
COLOR_MUTED = "8D99A6"
COLOR_HEADER_BG = "4E5B6E"
COLOR_BORDER = "D1D8DD"
COLOR_BAND = "F7F9FA"
COLOR_TOTAL_BG = "EBEFF2"

THIN = Side(style="thin", color=COLOR_BORDER)
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BORDER = Border(
	left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color=COLOR_HEADER_BG)
)


def clean(value):
	"""openpyxl refuses control characters."""
	if isinstance(value, str):
		value = ILLEGAL_CHARACTERS_RE.sub("", value)
	return value


def strip_html(value):
	from frappe.utils.xlsxutils import handle_html

	try:
		return handle_html(value)
	except Exception:
		return frappe.utils.strip_html(value)


def get_date_format():
	fmt = frappe.db.get_single_value("System Settings", "date_format") or "yyyy-mm-dd"
	return fmt.upper()


def get_time_format():
	fmt = frappe.db.get_single_value("System Settings", "time_format") or "HH:mm:ss"
	# frappe stores HH:mm:ss / HH:mm - excel wants the same letters uppercased
	return fmt.replace("mm", "MM").upper()


def number_format(df):
	"""Excel number format string for a numeric docfield."""
	fieldtype = df.get("fieldtype")

	if fieldtype == "Int":
		return "#,##0"

	if fieldtype == "Percent":
		return "0.00%"

	precision = cint(df.get("precision"))
	if not precision:
		if fieldtype == "Currency":
			precision = cint(frappe.db.get_default("currency_precision")) or 2
		else:
			precision = cint(frappe.db.get_default("float_precision")) or 2

	return "#,##0" + ("." + "0" * precision if precision else "")


def cell_value(value, df):
	"""Convert a stored value into something Excel understands natively."""
	fieldtype = df.get("fieldtype")

	if value in (None, ""):
		# 0 is meaningful for numbers, blank is meaningful for everything else
		return 0 if fieldtype in NUMERIC_TYPES else ""

	if fieldtype == "Check":
		return _("Yes") if cint(value) else _("No")

	if fieldtype == "Percent":
		# stored as 20 for 20%, but the 0.00% format multiplies by 100
		return flt(value) / 100.0

	if fieldtype in ("Currency", "Float"):
		return flt(value)

	if fieldtype == "Int":
		return cint(value)

	if fieldtype == "Date":
		try:
			return getdate(value)
		except Exception:
			return clean(str(value))

	if fieldtype == "Datetime":
		try:
			dt = get_datetime(value)
			# excel has no timezone concept, drop microseconds for a clean cell
			return dt.replace(microsecond=0)
		except Exception:
			return clean(str(value))

	if fieldtype == "Time":
		try:
			if isinstance(value, datetime.timedelta):
				total = int(value.total_seconds())
				return datetime.time(total // 3600 % 24, total // 60 % 60, total % 60)
			return get_datetime("2000-01-01 " + str(value)).time()
		except Exception:
			return clean(str(value))

	if fieldtype in HTML_TYPES:
		return clean(strip_html(frappe.as_unicode(value)))

	return clean(frappe.as_unicode(value))


def display_width(value, df):
	"""Rough rendered width of a value, used to size columns."""
	fieldtype = df.get("fieldtype")

	if isinstance(value, (datetime.datetime,)):
		return len(get_date_format()) + len(get_time_format()) + 1
	if isinstance(value, datetime.date):
		return len(get_date_format())
	if isinstance(value, datetime.time):
		return len(get_time_format())
	if isinstance(value, (int, float)) and fieldtype in NUMERIC_TYPES:
		# thousands separators and decimals add characters the raw repr lacks
		return len("{:,.2f}".format(value)) + 1

	text = frappe.as_unicode(value or "")
	if not text:
		return 0
	# multi-line text wraps, so only the longest line matters
	return max(len(line) for line in text.split("\n"))


def build_workbook(columns, rows, title, subtitle=None, add_totals=False, sheet_name=None):
	"""Return a BytesIO xlsx.

	:param columns: list of docfield-ish dicts, each needs `fieldname`, `label`, `fieldtype`
	:param rows: list of dicts keyed by fieldname
	:param title: heading printed in A1
	:param subtitle: small grey line under the title
	:param add_totals: append a SUM row for Currency/Float/Int columns
	"""
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = safe_sheet_name(sheet_name or title)

	last_col = len(columns)
	last_letter = get_column_letter(last_col)

	# ---- title block --------------------------------------------------
	ws["A1"] = clean(title)
	ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=COLOR_TEXT)
	ws["A2"] = clean(subtitle or "")
	ws["A2"].font = Font(name="Calibri", size=9, italic=True, color=COLOR_MUTED)
	if last_col > 1:
		ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
		ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
	ws.row_dimensions[1].height = 22
	ws.row_dimensions[3].height = 6

	# ---- header row ---------------------------------------------------
	header_row = 4
	header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
	widths = []

	for i, df in enumerate(columns, start=1):
		label = _(df.get("label") or df.get("fieldname"))
		cell = ws.cell(row=header_row, column=i, value=clean(label))
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = CELL_BORDER
		widths.append(len(label))

	ws.row_dimensions[header_row].height = 24

	# ---- data ---------------------------------------------------------
	band_fill = PatternFill("solid", fgColor=COLOR_BAND)
	date_fmt = get_date_format()
	time_fmt = get_time_format()

	for r, row in enumerate(rows):
		excel_row = header_row + 1 + r
		banded = r % 2 == 1

		for i, df in enumerate(columns, start=1):
			fieldtype = df.get("fieldtype")
			value = cell_value(row.get(df.get("fieldname")), df)
			cell = ws.cell(row=excel_row, column=i, value=value)
			cell.border = CELL_BORDER
			cell.font = Font(name="Calibri", size=10, color=COLOR_TEXT)

			if banded:
				cell.fill = band_fill

			if fieldtype in NUMERIC_TYPES:
				cell.number_format = number_format(df)
				cell.alignment = Alignment(horizontal="right", vertical="top")
			elif fieldtype == "Date":
				cell.number_format = date_fmt
				cell.alignment = Alignment(horizontal="center", vertical="top")
			elif fieldtype == "Datetime":
				cell.number_format = date_fmt + " " + time_fmt
				cell.alignment = Alignment(horizontal="center", vertical="top")
			elif fieldtype == "Time":
				cell.number_format = time_fmt
				cell.alignment = Alignment(horizontal="center", vertical="top")
			elif fieldtype == "Check":
				cell.alignment = Alignment(horizontal="center", vertical="top")
			else:
				cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

			widths[i - 1] = max(widths[i - 1], display_width(value, df))

	last_data_row = header_row + len(rows)

	# ---- totals -------------------------------------------------------
	if add_totals and rows:
		total_row = last_data_row + 1
		ws.cell(row=total_row, column=1, value=_("Total"))

		for i, df in enumerate(columns, start=1):
			cell = ws.cell(row=total_row, column=i)
			cell.font = Font(name="Calibri", bold=True, size=10, color=COLOR_TEXT)
			cell.fill = PatternFill("solid", fgColor=COLOR_TOTAL_BG)
			cell.border = TOTAL_BORDER

			if df.get("fieldtype") in SUMMABLE_TYPES and df.get("fieldname") != "idx":
				letter = get_column_letter(i)
				cell.value = "=SUM({0}{1}:{0}{2})".format(letter, header_row + 1, last_data_row)
				cell.number_format = number_format(df)
				cell.alignment = Alignment(horizontal="right")
			elif i == 1:
				cell.alignment = Alignment(horizontal="left")

	# ---- finishing ----------------------------------------------------
	for i, width in enumerate(widths, start=1):
		ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 55)

	ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
	if rows:
		ws.auto_filter.ref = "A{0}:{1}{2}".format(header_row, last_letter, last_data_row)

	ws.sheet_view.showGridLines = False
	ws.page_setup.orientation = "landscape" if last_col > 6 else "portrait"
	ws.print_title_rows = "{0}:{0}".format(header_row)

	out = BytesIO()
	wb.save(out)
	return out


def safe_sheet_name(name):
	name = INVALID_SHEET_CHARS_RE.sub(" ", frappe.as_unicode(name or "Sheet1")).strip()
	return (name or "Sheet1")[:31]


def safe_filename(name):
	name = re.sub(r"[^\w\-. ]+", "_", frappe.as_unicode(name or "export")).strip()
	return (name or "export")[:120]
